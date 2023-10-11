import requests
from urllib.parse import urlencode
import openpyxl

# Replace these placeholders with your actual eBulkSMS API credentials
base_url = "https://api.ebulksms.com:8080/sendsms"
username = "yourmail@anymail.com"
apikey = "XXXXXb8d09da177519bd9cXXXXXXXXXXXXX"
sender_name = "TEST"

# Load the Excel file
excel_file = "test.xlsx"
workbook = openpyxl.load_workbook(excel_file)
sheet = workbook.active

# Iterate through rows in the sheet
for row in sheet.iter_rows(min_row=2, values_only=True):
    phone_number, account_balance = row
    # Customize the message as needed
    message_text = (
        f"Hello, your account balance wih BSD Cooperative is NGN{account_balance}. For more, contact 0802345701."
    )
    # \u20A6 for naira symbol
    # Construct the payload for the request
    payload = {
        "username": username,
        "apikey": apikey,
        "sender": sender_name,
        "messagetext": message_text,
        "flash": 0,
        "recipients": phone_number,
    }

    try:
        # Send a GET request to the API URL with URL-encoded payload
        api_url = f"{base_url}?{urlencode(payload)}"
        response = requests.get(api_url)

        # Check if the request was successful (HTTP status code 200)
        if response.status_code == 200:
            # The response text contains the result of the SMS sending operation
            result = response.text
            print(f"SMS sent to {phone_number}: {result}")
        else:
            print(
                f"Request failed for {phone_number} with status code {response.status_code}"
            )

    except requests.exceptions.RequestException as e:
        print(f"Request error for {phone_number}: {e}")

# Close the Excel file
workbook.close()
